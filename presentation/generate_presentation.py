from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "presentation"
ASSET_DIR = ROOT / "images"
OUTPUT = OUT_DIR / "ai_document_processing_workflow.pptx"

NAVY = RGBColor(12, 23, 40)
NAVY_2 = RGBColor(20, 36, 59)
TEAL = RGBColor(37, 191, 174)
CYAN = RGBColor(111, 222, 216)
ORANGE = RGBColor(245, 163, 72)
WHITE = RGBColor(245, 248, 250)
MUTED = RGBColor(168, 184, 199)
DARK = RGBColor(24, 34, 48)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, color, radius=False, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line or color
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False, font="Aptos", align=None, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align or PP_ALIGN.LEFT
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def bullet_list(slide, items, x, y, w, h, size=16, color=WHITE, gap=0.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.04)
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(gap * 10)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p._p.get_or_add_pPr().insert(0, p._p._new_buChar())
    return box


def base(title, kicker=None, number=None):
    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    rect(slide, 0, 0, 0.18, 7.5, TEAL)
    if number is not None:
        text(slide, f"{number:02d}", 12.25, 0.32, 0.55, 0.3, size=12, color=TEAL, bold=True, align=PP_ALIGN.RIGHT)
    if kicker:
        text(slide, kicker.upper(), 0.72, 0.42, 5, 0.25, size=10, color=TEAL, bold=True)
    text(slide, title, 0.72, 0.72, 11.2, 0.62, size=28, color=WHITE, bold=True)
    rect(slide, 0.72, 1.48, 1.0, 0.05, ORANGE)
    return slide


def card(slide, x, y, w, h, title, body, accent=TEAL, icon=None):
    rect(slide, x, y, w, h, NAVY_2, radius=True, line=RGBColor(42, 61, 83))
    rect(slide, x, y, 0.07, h, accent)
    if icon:
        text(slide, icon, x + 0.25, y + 0.25, 0.45, 0.4, size=22, color=accent, bold=True)
        tx = x + 0.78
        tw = w - 1.0
    else:
        tx = x + 0.28
        tw = w - 0.55
    text(slide, title, tx, y + 0.25, tw, 0.36, size=16, color=WHITE, bold=True)
    text(slide, body, tx, y + 0.78, tw, h - 0.98, size=12.5, color=MUTED)


def add_footer(slide, label="AI Document Processing"):
    text(slide, label, 0.72, 7.1, 4.0, 0.2, size=9, color=MUTED)
    text(slide, "GenAI Capstone Project", 9.6, 7.1, 2.95, 0.2, size=9, color=MUTED, align=PP_ALIGN.RIGHT)


# 1. Title
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
rect(slide, 0, 0, 0.25, 7.5, TEAL)
rect(slide, 8.65, 0, 4.68, 7.5, NAVY_2)
rect(slide, 9.45, 1.05, 2.55, 0.08, ORANGE)
text(slide, "GENAI CAPSTONE PROJECT", 0.9, 1.0, 4.7, 0.3, size=12, color=TEAL, bold=True)
text(slide, "AI-Powered\nDocument Processing\n& Business Workflow", 0.88, 1.6, 7.4, 2.35, size=31, color=WHITE, bold=True)
text(slide, "From customer complaint documents to validated cases, customer communication, management summaries, and consolidated reports.", 0.92, 4.45, 6.85, 0.8, size=16, color=MUTED)
for idx, label in enumerate(["TXT / PDF / DOCX", "LLM + Pydantic", "FastAPI + Streamlit"]):
    rect(slide, 0.92 + idx * 2.18, 5.75, 1.9, 0.42, NAVY_2, radius=True, line=RGBColor(42, 61, 83))
    text(slide, label, 1.02 + idx * 2.18, 5.86, 1.7, 0.18, size=9.5, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
text(slide, "A practical, testable GenAI workflow for customer case operations", 9.15, 5.65, 3.55, 0.75, size=17, color=WHITE, bold=True)
text(slide, "01", 12.25, 0.4, 0.5, 0.25, size=12, color=TEAL, bold=True, align=PP_ALIGN.RIGHT)

# 2. Problem and solution
slide = base("Turning unstructured complaints into operational decisions", "Business context", 2)
card(slide, 0.72, 1.95, 5.55, 3.75, "The challenge", "Complaints arrive as unstructured files. Manual review creates slow turnaround, inconsistent classification, repetitive writing, and limited visibility into failed cases.", ORANGE, "!")
card(slide, 7.0, 1.95, 5.55, 3.75, "The solution", "A modular pipeline parses each document, extracts a validated case, generates two business outputs, and records every result in structured reports.", TEAL, ">")
text(slide, "The result", 0.72, 6.12, 1.4, 0.25, size=11, color=ORANGE, bold=True)
text(slide, "Faster review  •  Consistent data  •  Traceable outcomes", 2.02, 6.08, 8.6, 0.35, size=20, color=WHITE, bold=True)
add_footer(slide)

# 3. What the system does
slide = base("One workflow, multiple business outputs", "Capabilities", 3)
items = [
    ("01", "Ingest", "TXT, PDF, and DOCX complaint files."),
    ("02", "Extract", "Convert unstructured text into ComplaintCase."),
    ("03", "Communicate", "Draft a grounded customer response email."),
    ("04", "Summarize", "Create an internal management summary."),
    ("05", "Report", "Persist JSON, TXT, CSV, and processing status."),
    ("06", "Deliver", "Expose the flow through API and browser UI."),
]
for i, (num, title, body) in enumerate(items):
    x = 0.72 + (i % 3) * 4.15
    y = 1.95 + (i // 3) * 2.05
    rect(slide, x, y, 3.6, 1.55, NAVY_2, radius=True, line=RGBColor(42, 61, 83))
    text(slide, num, x + 0.25, y + 0.22, 0.55, 0.3, size=12, color=ORANGE, bold=True)
    text(slide, title, x + 0.95, y + 0.2, 2.2, 0.3, size=16, color=WHITE, bold=True)
    text(slide, body, x + 0.95, y + 0.7, 2.35, 0.58, size=12, color=MUTED)
add_footer(slide)

# 4. Architecture image
slide = base("A layered architecture built for change", "System design", 4)
image = ASSET_DIR / "architecture.png"
if image.exists():
    slide.shapes.add_picture(str(image), Inches(0.85), Inches(1.78), width=Inches(7.65), height=Inches(4.85))
card(slide, 8.8, 1.95, 3.55, 1.25, "Application layer", "Batch processor and workflow router coordinate the business process.", TEAL)
card(slide, 8.8, 3.4, 3.55, 1.25, "AI layer", "LLM client centralizes tiers, retries, fallbacks, and structured responses.", ORANGE)
card(slide, 8.8, 4.85, 3.55, 1.25, "Delivery layer", "FastAPI and Streamlit make the workflow usable by systems and people.", CYAN)
add_footer(slide)

# 5. End-to-end workflow
slide = base("Sequential extraction, parallel generation", "Workflow", 5)
steps = [("01", "Parse", "Load and group document content"), ("02", "Extract", "Create ComplaintCase with the LLM"), ("03", "Validate", "Enforce the Pydantic contract"), ("04", "Generate", "Email + summary run independently"), ("05", "Persist", "Save artifacts and report status")]
for i, (num, title, body) in enumerate(steps):
    x = 0.72 + i * 2.48
    rect(slide, x, 2.25, 2.05, 2.15, NAVY_2, radius=True, line=RGBColor(42, 61, 83))
    text(slide, num, x + 0.2, 2.52, 0.45, 0.25, size=12, color=ORANGE, bold=True)
    text(slide, title, x + 0.2, 3.0, 1.6, 0.3, size=16, color=WHITE, bold=True)
    text(slide, body, x + 0.2, 3.55, 1.58, 0.55, size=11.5, color=MUTED)
    if i < 4:
        text(slide, ">", x + 2.13, 3.08, 0.3, 0.3, size=21, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
rect(slide, 4.25, 5.05, 4.85, 0.72, NAVY_2, radius=True, line=TEAL)
text(slide, "Parallel branch: customer email  +  management summary", 4.5, 5.27, 4.35, 0.25, size=15, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
add_footer(slide)

# 6. Structured extraction
slide = base("The LLM is constrained by a business contract", "Structured intelligence", 6)
text(slide, "Unstructured complaint", 0.85, 2.0, 2.1, 0.3, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
rect(slide, 0.85, 2.55, 2.1, 1.65, NAVY_2, radius=True, line=RGBColor(42, 61, 83))
text(slide, "Customer reports\nissue, action, status,\nand supporting evidence", 1.05, 2.92, 1.7, 0.75, size=14, color=MUTED, align=PP_ALIGN.CENTER)
text(slide, ">", 3.35, 3.05, 0.45, 0.4, size=26, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
text(slide, "LLM extraction", 4.0, 2.0, 2.2, 0.3, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
rect(slide, 4.0, 2.55, 2.2, 1.65, NAVY_2, radius=True, line=RGBColor(42, 61, 83))
text(slide, "Grounded prompts\n+ model tier\n+ fallback", 4.25, 2.92, 1.7, 0.75, size=14, color=MUTED, align=PP_ALIGN.CENTER)
text(slide, ">", 6.65, 3.05, 0.45, 0.4, size=26, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
text(slide, "Validated case", 7.3, 2.0, 2.2, 0.3, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
rect(slide, 7.3, 2.55, 2.2, 1.65, NAVY_2, radius=True, line=RGBColor(42, 61, 83))
text(slide, "ComplaintCase\nPydantic schema\nstrict fields", 7.55, 2.92, 1.7, 0.75, size=14, color=MUTED, align=PP_ALIGN.CENTER)
card(slide, 10.15, 2.05, 2.2, 2.35, "Why it matters", "The downstream workflows receive predictable data instead of an unvalidated raw response.", ORANGE)
add_footer(slide)

# 7. LLM strategy
slide = base("Reliability is designed into the LLM boundary", "LLM strategy", 7)
card(slide, 0.72, 1.95, 3.65, 3.75, "Model tiers", "BASIC handles email and summary generation. GENERIC handles structured extraction. COMPLEX is available for future workflows.", TEAL, "1")
card(slide, 4.85, 1.95, 3.65, 3.75, "Retry + fallback", "Each tier has a primary and fallback model. Requests use configured timeout and retry policies before failure is reported.", ORANGE, "2")
card(slide, 8.98, 1.95, 3.65, 3.75, "Grounded prompts", "Prompts require source-only facts and prohibit invented customer details, actions, dates, resolutions, and promises.", CYAN, "3")
text(slide, "Failure is visible: it becomes a logged, document-level result in the final report.", 1.6, 6.25, 10.0, 0.35, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_footer(slide)

# 8. Interfaces with screenshots
slide = base("Built for both people and systems", "Delivery", 8)
ui = ASSET_DIR / "ui_flow.png"
api = ASSET_DIR / "api_doc.png"
if ui.exists():
    slide.shapes.add_picture(str(ui), Inches(0.78), Inches(1.85), width=Inches(5.75), height=Inches(4.75))
if api.exists():
    slide.shapes.add_picture(str(api), Inches(6.85), Inches(1.85), width=Inches(5.75), height=Inches(4.75))
text(slide, "Streamlit UI", 2.55, 6.67, 2.0, 0.25, size=14, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
text(slide, "FastAPI + Swagger", 8.65, 6.67, 2.2, 0.25, size=14, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
add_footer(slide)

# 9. Outputs
slide = base("Every run leaves a useful audit trail", "Outputs", 9)
card(slide, 0.72, 1.95, 3.6, 3.8, "Structured data", "One JSON file per source document with customer details, category, issue, resolution, escalation, evidence, and case status.", TEAL, "JSON")
card(slide, 4.85, 1.95, 3.6, 3.8, "Human-readable outputs", "A customer-facing email and an internal management summary are saved as separate text artifacts.", ORANGE, "TXT")
card(slide, 8.98, 1.95, 3.6, 3.8, "Batch report", "The final CSV records source files, success or failure, errors, and paths to generated outputs.", CYAN, "CSV")
text(slide, "Local batch: output/     API runs: output/runs/<run_id>/", 1.45, 6.3, 10.5, 0.3, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_footer(slide)

# 10. API
slide = base("A small API surface with a complete workflow behind it", "API", 10)
endpoints = [("GET", "/health", "Service health check"), ("POST", "/api/v1/process", "Upload and process documents"), ("GET", "/api/v1/report/{run_id}", "Download the final CSV report")]
for i, (method, route, desc) in enumerate(endpoints):
    y = 1.95 + i * 1.35
    rect(slide, 0.9, y, 1.1, 0.55, TEAL if method == "POST" else NAVY_2, radius=True, line=TEAL)
    text(slide, method, 1.0, y + 0.16, 0.9, 0.2, size=12, color=NAVY if method == "POST" else TEAL, bold=True, align=PP_ALIGN.CENTER)
    text(slide, route, 2.35, y + 0.08, 4.4, 0.3, size=18, color=WHITE, bold=True)
    text(slide, desc, 7.3, y + 0.1, 4.3, 0.3, size=14, color=MUTED)
rect(slide, 1.0, 6.0, 11.2, 0.62, NAVY_2, radius=True, line=RGBColor(42, 61, 83))
text(slide, "Upload  →  Run ID  →  Processing statistics  →  Downloadable report", 1.25, 6.2, 10.7, 0.2, size=16, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
add_footer(slide)

# 11. Testing and operations
slide = base("Designed to be tested without spending on live calls", "Quality and operations", 11)
card(slide, 0.72, 1.95, 3.7, 3.75, "Automated tests", "Pytest coverage spans parsing, schemas, LLM behavior, workflows, batch processing, output persistence, and FastAPI endpoints.", TEAL, "✓")
card(slide, 4.82, 1.95, 3.7, 3.75, "Mocked LLM calls", "Tests remain deterministic and generally do not require an API key or provider calls.", ORANGE, "◇")
card(slide, 8.92, 1.95, 3.7, 3.75, "Operational visibility", "Console and rotating file logs capture document-level outcomes, API run IDs, failures, and model activity.", CYAN, "≡")
text(slide, "Run: python -m pytest", 4.2, 6.25, 4.9, 0.32, size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_footer(slide)

# 12. Demo path / closing
slide = base("A clear path from upload to decision", "Demo flow", 12)
steps = ["Upload complaints", "Watch processing", "Inspect report", "Review email", "Review summary"]
for i, label in enumerate(steps):
    x = 0.8 + i * 2.42
    rect(slide, x, 2.05, 1.9, 1.1, NAVY_2, radius=True, line=TEAL)
    text(slide, f"{i + 1}", x + 0.72, 2.2, 0.45, 0.3, size=18, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    text(slide, label, x + 0.15, 2.72, 1.6, 0.22, size=11.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    if i < 4:
        text(slide, ">", x + 2.02, 2.42, 0.25, 0.3, size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
text(slide, "One complaint document becomes a validated case and three useful outputs.", 1.2, 4.25, 10.95, 0.6, size=24, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(slide, "AI-assisted processing with validation, traceability, and human review built into the workflow.", 1.8, 5.2, 9.75, 0.4, size=16, color=MUTED, align=PP_ALIGN.CENTER)
rect(slide, 4.45, 6.18, 4.4, 0.08, ORANGE)
add_footer(slide, "AI Document Processing  •  Thank you")

prs.save(OUTPUT)
print(OUTPUT)
